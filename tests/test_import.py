import io

import pytest
from docx import Document as DocxDocument

from app.content import FileImportError, import_file


def make_docx() -> bytes:
    d = DocxDocument()
    d.add_heading("Launch plan", level=1)
    p = d.add_paragraph()
    p.add_run("Bold ").bold = True
    p.add_run("and italic").italic = True
    d.add_paragraph("first item", style="List Bullet")
    d.add_paragraph("second item", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_keeps_structure():
    title, html = import_file("launch_plan.docx", make_docx())
    assert title == "launch plan"
    assert "<h1>Launch plan</h1>" in html
    assert "<strong>Bold </strong>" in html and "<em>and italic</em>" in html
    assert "<ul><li>first item</li><li>second item</li></ul>" in html


def test_markdown_import_is_sanitized():
    src = b"# Title\n\n- **one**\n- two\n\n<script>alert(1)</script>\n\n#### deep"
    _, html = import_file("notes.md", src)
    assert "<h1>Title</h1>" in html
    assert "<li><strong>one</strong></li>" in html
    assert "<h3>deep</h3>" in html  # h4-h6 folded into h3
    assert "script" not in html


def test_txt_import_escapes_and_splits_paragraphs():
    _, html = import_file("a.txt", "line 1\nline <2>\n\npara 2".encode())
    assert html == "<p>line 1<br>line &lt;2&gt;</p><p>para 2</p>"


@pytest.mark.parametrize("name,data,msg", [
    ("virus.exe", b"MZ", "Unsupported file type"),
    ("empty.txt", b"", "empty"),
    ("broken.docx", b"not a zip", "valid Word document"),
    ("latin.txt", "caf\xe9".encode("latin-1"), "UTF-8"),
])
def test_bad_files_give_clear_errors(name, data, msg):
    with pytest.raises(FileImportError, match=msg):
        import_file(name, data)


def test_upload_endpoint_creates_owned_document(alice):
    res = alice.post(
        "/documents/import",
        data={"file": (io.BytesIO(make_docx()), "launch_plan.docx")},
        content_type="multipart/form-data",
    )
    assert res.status_code == 302
    doc_id = int(res.headers["Location"].split("/")[-1])
    doc = alice.get(f"/api/documents/{doc_id}").json
    assert doc["role"] == "owner" and doc["title"] == "launch plan"
    assert "<h1>Launch plan</h1>" in doc["content"]


def test_upload_rejects_unsupported_type(alice):
    res = alice.post(
        "/documents/import",
        data={"file": (io.BytesIO(b"%PDF"), "scan.pdf")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"Unsupported file type" in res.data
